#!/usr/bin/env python3
"""Optional editable export. Build main.aux and architecture.png first; PDF is canonical."""
from pathlib import Path
import re,subprocess,os
root=Path(__file__).resolve().parents[1];report=root/'docs/report';aux=(report/'output/main.aux').read_text();labels=dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',aux))
body='\n'.join((report/'sections'/n).read_text() for n in ['01-operating-systems.tex','02-website.tex','03-security.tex','04-handover-wbs.tex'])
scope=(report/'main.tex').read_text();scope=scope[scope.index('\\chapter*{Phạm vi'):scope.index('\\setcounter{tocdepth}')];body=scope+'\n'+body
body=body.replace('\\input{sections/plugin-table}',(report/'sections/plugin-table.tex').read_text())
body=re.sub(r'\\ref\{([^}]+)\}',lambda m:labels.get(m[1],'[cần cập nhật]'),body)
body=re.sub(r'\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}',r'\\includegraphics[width=0.8\\linewidth]{figures/architecture.png}',body,flags=re.S)
def shot(m):
 name,caption,label=m.group(1),m.group(2),m.group(3)
 return '\\begin{figure}\n\\includegraphics[width=0.9\\linewidth]{figures/'+name+'}\n\\caption{Hình '+labels[label]+'. '+caption+'}\n\\end{figure}'
body=re.sub(r'\\shot(?:\[[^\]]*\])?\{([^}]+)\}\{((?:[^{}]|\{[^{}]*\})*)\}\{([^}]+)\}',shot,body)
def source(m):
 name,start,end=m.group(1),int(m.group(2)),int(m.group(3));text='\n'.join((root/name).read_text().splitlines()[start-1:end]);return '\\begin{verbatim}\n'+text+'\n\\end{verbatim}'
body=re.sub(r'\\sourcefile\{([^}]+)\}\{(\d+)\}\{(\d+)\}',source,body)
body=body.replace('\\code{','\\texttt{')
# Pandoc supports standard p columns; no custom LaTeX column macros in the DOCX input.
body=body.replace('P{','p{')
columns={'lYY':'p{3cm}p{5.75cm}p{5.75cm}','lYYY':'p{2.5cm}p{3.7cm}p{3.7cm}p{3.7cm}','YYY':'p{5cm}p{5cm}p{4.5cm}','YYl':'p{5.75cm}p{5.75cm}p{3cm}'}
for original,standard in columns.items():
    body=body.replace('\\begin{tabularx}{\\linewidth}{'+original+'}', '\\begin{tabular}{'+standard+'}')
body=body.replace('\\end{tabularx}', '\\end{tabular}')
body=re.sub(r'\\begin\{minipage\}(?:\[[^]]*\])?\{[^}]*\}', '', body).replace('\\end{minipage}', '').replace('\\captionof{figure}', '\\caption')
def numbered_caption(match):
    caption, label = match.group(1), match.group(2)
    kind = 'Bảng' if label.startswith('tab:') else 'Hình'
    return '\\caption{'+kind+' '+labels[label]+'. '+caption+'}'
body=re.sub(r'\\caption(?:\[[^\]]*\])?\{((?:[^{}]|\{[^{}]*\})*)\}\\label\{([^}]+)\}',numbered_caption,body)
tex='\\documentclass{report}\n\\begin{document}\n'+body+'\n\\bibliography{references}\n\\end{document}\n'
from tempfile import TemporaryDirectory
with TemporaryDirectory(prefix='lk-docx-') as directory:
    expanded=Path(directory)/'expanded.tex';expanded.write_text(tex)
    subprocess.run([os.environ.get('PANDOC','pandoc'),str(expanded),'--from=latex','--to=docx','--standalone','--citeproc','--bibliography=references.bib','--resource-path=.:figures','-o','output/lang-kinh-editable.docx'],cwd=report,check=True)
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
p=report/'output/lang-kinh-editable.docx';d=Document(p)
for sec in d.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
for style in ['Normal','Body Text']:
 d.styles[style].font.name='Liberation Serif';d.styles[style].font.size=Pt(11)
 d.styles[style].paragraph_format.line_spacing=1.12
nodes=[]
for text,bold,size in [('TRƯỜNG: ______________________________',False,13),('KHOA: _________________________________',False,13),('',False,13),('BÁO CÁO HỌC PHẦN',True,18),('Môn học: ______________________________',False,13),('',False,13),('THƯ VIỆN ẢNH SỐ LĂNG KÍNH',True,20),('Nhóm 1',True,16),('Lớp: __________________________________',False,13),('Giảng viên: ____________________________',False,13)]:
 par=d.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=par.add_run(text);r.bold=bold;r.font.size=Pt(size);nodes.append(par._p)
t=d.add_table(rows=1,cols=3)
for cell,text in zip(t.rows[0].cells,['Thành viên','Họ và tên','MSSV']):cell.text=text
for i in range(1,6):
 cells=t.add_row().cells;cells[0].text='TV'+str(i);cells[1].text='________________________';cells[2].text='________________'
for row in t.rows:
 row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
nodes.append(t._tbl)
par=d.add_paragraph('Tháng 10 năm 2026');par.alignment=WD_ALIGN_PARAGRAPH.CENTER;nodes.append(par._p)
par=d.add_paragraph();par.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE);nodes.append(par._p)
for i,node in enumerate(nodes):d._element.body.insert(i,node)
d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title='Lăng Kính — Nhóm 1';d.core_properties.subject='Bản chỉnh sửa bổ sung; PDF LaTeX là bản trình bày chính'
d.save(p)
print('Editable DOCX exported with explicit cover, actual code excerpts, diagram and static references from current PDF.')